"""ATS Resume Checker - Streamlit + Groq.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description, and get:
  * an ATS compatibility score (0-100) with a category breakdown
  * strengths, problems, missing keywords
  * concrete improvements and rewritten bullet examples
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "llama-3.3-70b-versatile"
MODEL_CHOICES = ["llama-3.3-70b-versatile", "llama-3.1-8b-instant"]
MAX_RESUME_CHARS = 12000  # keeps the prompt well inside the model context
MAX_JD_CHARS = 6000
MIN_TEXT_CHARS = 100  # below this the file is probably a scanned image

CATEGORY_KEYS = [
    ("formatting", "Formatting & Structure"),
    ("keywords", "Keywords & Skills"),
    ("experience", "Experience & Achievements"),
    ("education", "Education & Certifications"),
    ("readability", "Readability & Clarity"),
]

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. You evaluate how well a resume would be parsed and \
ranked by an ATS and how strong it is for a recruiter.

Rules:
- Base every statement ONLY on the resume text provided. Never invent experience.
- Be honest and specific. Do not inflate scores.
- Scores are integers from 0 to 100.
- Respond with a single valid JSON object and nothing else."""

USER_PROMPT_TEMPLATE = """Analyse the resume below{jd_clause}.

Return JSON with EXACTLY this structure:
{{
  "overall_score": <int 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {{
    "formatting": <int 0-100>,
    "keywords": <int 0-100>,
    "experience": <int 0-100>,
    "education": <int 0-100>,
    "readability": <int 0-100>
  }},
  "strengths": ["<short point>", "..."],
  "issues": ["<specific problem found in this resume>", "..."],
  "missing_keywords": ["<keyword or skill>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "suggestion": "<actionable fix>"}}
  ],
  "bullet_rewrites": [
    {{"original": "<line copied from the resume>", "improved": "<stronger version with action verb and measurable result placeholder>"}}
  ]
}}

Guidelines:
- 3-6 strengths, 3-8 issues, up to 12 missing keywords, 5-8 improvements, 2-4 bullet rewrites.
- {keyword_rule}
- In "improved" bullets, use placeholders like [X%] or [number] where a metric is needed; do not make up facts.

{jd_block}RESUME TEXT:
\"\"\"
{resume_text}
\"\"\""""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = file_name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                unlocked = bool(reader.decrypt(""))
            except Exception:
                unlocked = False
            if not unlocked:
                raise ValueError("This PDF is password protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Quick rule-based checks (no AI needed)
# ----------------------------------------------------------------------------
def basic_checks(text: str) -> list:
    """Return a list of (passed, message) tuples for simple ATS checks."""
    lowered = text.lower()
    words = len(text.split())
    checks = [
        (
            bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
            "Email address found",
        ),
        (
            bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)),
            "Phone number found",
        ),
        (
            any(k in lowered for k in ("experience", "employment", "work history")),
            "Work experience section detected",
        ),
        (
            any(k in lowered for k in ("education", "university", "college", "bachelor", "master")),
            "Education section detected",
        ),
        (
            "skills" in lowered,
            "Skills section detected",
        ),
        (
            250 <= words <= 1200,
            f"Length looks reasonable ({words} words; ideal is roughly 250-1200)",
        ),
    ]
    return checks


# ----------------------------------------------------------------------------
# Groq call + parsing
# ----------------------------------------------------------------------------
def get_api_key() -> str:
    """Read the key from Streamlit secrets, then environment variables."""
    try:
        key = st.secrets.get("GROQ_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GROQ_API_KEY", "")


def build_prompt(resume_text: str, job_description: str) -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    if job_description.strip():
        jd_clause = " and the target job description"
        keyword_rule = (
            "For missing_keywords, compare the resume against the job description "
            "and list important terms from the job description that the resume lacks."
        )
        jd_block = (
            'JOB DESCRIPTION:\n"""\n' + job_description.strip()[:MAX_JD_CHARS] + '\n"""\n\n'
        )
    else:
        jd_clause = ""
        keyword_rule = (
            "For missing_keywords, list common industry keywords and skills that "
            "are typical for the candidate's apparent field but absent from the resume."
        )
        jd_block = ""
    return USER_PROMPT_TEMPLATE.format(
        jd_clause=jd_clause,
        keyword_rule=keyword_rule,
        jd_block=jd_block,
        resume_text=resume_text,
    )


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences or extra text."""
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise


def _score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalise_result(data: dict) -> dict:
    """Make sure every field exists and has the right type."""
    if not isinstance(data, dict):
        raise ValueError("Model did not return a JSON object.")

    cats = data.get("category_scores")
    cats = cats if isinstance(cats, dict) else {}

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict) and item.get("suggestion"):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).title(),
                    "suggestion": str(item["suggestion"]).strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append({"priority": "Medium", "suggestion": item.strip()})

    rewrites = []
    for item in data.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append(
                {"original": str(item["original"]).strip(), "improved": str(item["improved"]).strip()}
            )

    return {
        "overall_score": _score(data.get("overall_score")),
        "summary": str(data.get("summary", "")).strip(),
        "category_scores": {key: _score(cats.get(key)) for key, _ in CATEGORY_KEYS},
        "strengths": _str_list(data.get("strengths")),
        "issues": _str_list(data.get("issues")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "improvements": improvements,
        "bullet_rewrites": rewrites,
    }


def analyse_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    client = Groq(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": prompt},
    ]

    try:
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            temperature=0.2,
            max_tokens=2500,
            response_format={"type": "json_object"},
        )
    except Exception:
        # Some models do not support JSON mode - retry without it.
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            temperature=0.2,
            max_tokens=2500,
        )

    raw = response.choices[0].message.content or ""
    return normalise_result(parse_json_response(raw))


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict, resume_text: str) -> None:
    overall = result["overall_score"]

    st.subheader("Your ATS score")
    col1, col2 = st.columns([1, 3])
    with col1:
        st.metric("Overall", f"{overall}/100", score_label(overall), delta_color="off")
    with col2:
        st.progress(overall / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(len(CATEGORY_KEYS))
    for col, (key, label) in zip(cols, CATEGORY_KEYS):
        with col:
            value = result["category_scores"][key]
            st.metric(label, f"{value}/100")
            st.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for item in result["strengths"] or ["No strengths reported."]:
            st.markdown(f"- {item}")
    with right:
        st.subheader("Issues found")
        for item in result["issues"] or ["No issues reported."]:
            st.markdown(f"- {item}")

    st.subheader("Missing keywords")
    if result["missing_keywords"]:
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
    else:
        st.write("No missing keywords reported.")

    st.subheader("How to improve your resume")
    order = {"High": 0, "Medium": 1, "Low": 2}
    for item in sorted(result["improvements"], key=lambda i: order.get(i["priority"], 1)):
        st.markdown(f"**[{item['priority']}]** {item['suggestion']}")
    if not result["improvements"]:
        st.write("No improvements reported.")

    if result["bullet_rewrites"]:
        st.subheader("Example bullet rewrites")
        for item in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {item['original']}")
            st.markdown(f"**After:** {item['improved']}")
            st.divider()

    st.subheader("Quick checks")
    for passed, message in basic_checks(resume_text):
        st.markdown(f"{'✅' if passed else '❌'} {message}")

    report = json.dumps(result, indent=2)
    st.download_button(
        "Download report (JSON)",
        data=report,
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it. Powered by Groq.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if not api_key:
            api_key = st.text_input("Groq API key", type="password", help="Get one free at console.groq.com/keys")
        else:
            st.success("Groq API key loaded.")

        env_model = os.environ.get("GROQ_MODEL", "")
        choices = MODEL_CHOICES.copy()
        if env_model and env_model not in choices:
            choices.insert(0, env_model)
        model = st.selectbox("Model", choices, index=0)
        st.markdown("---")
        st.caption("Your resume is sent to Groq for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=160,
        placeholder="Paste the job description here for a targeted keyword match...",
    )

    if st.button("Analyse resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Groq API key in the sidebar (or in Streamlit secrets).")
            return

        try:
            resume_text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(resume_text) < MIN_TEXT_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image or "
                "uses text inside images, an ATS cannot read it either - export a text-based "
                "PDF or DOCX and try again."
            )
            return

        with st.spinner("Analysing your resume..."):
            try:
                result = analyse_resume(api_key, model, resume_text, job_description)
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
                return

        render_results(result, resume_text)


if __name__ == "__main__":
    main()
